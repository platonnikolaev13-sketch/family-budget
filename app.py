import os
import io
import json
import sqlite3
import secrets
from datetime import datetime, date
from functools import wraps

from flask import (Flask, request, jsonify, render_template, session,
                   send_file, g)
from flask_socketio import SocketIO, emit, join_room
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', secrets.token_hex(32))
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['PERMANENT_SESSION_LIFETIME'] = 60 * 60 * 24 * 365  # год

DB_PATH = os.environ.get('DB_PATH', 'budget.db')

# async_mode='threading' — работает на Render free без Redis
socketio = SocketIO(app, cors_allowed_origins='*',
                    async_mode='threading', logger=False, engineio_logger=False)


# ==================== DATABASE ====================
def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute('PRAGMA foreign_keys = ON')
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop('db', None)
    if db is not None:
        db.close()


def init_db():
    con = sqlite3.connect(DB_PATH)
    con.executescript('''
        CREATE TABLE IF NOT EXISTS families (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            invite_code TEXT UNIQUE NOT NULL,
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            display_name TEXT NOT NULL,
            family_id INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY (family_id) REFERENCES families(id)
        );

        CREATE TABLE IF NOT EXISTS expenses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            family_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            type TEXT NOT NULL CHECK(type IN ('expense','income')),
            name TEXT NOT NULL,
            amount REAL NOT NULL CHECK(amount > 0),
            category TEXT NOT NULL,
            date TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (family_id) REFERENCES families(id),
            FOREIGN KEY (user_id) REFERENCES users(id)
        );
        CREATE INDEX IF NOT EXISTS idx_exp_family_date
            ON expenses(family_id, date);

        CREATE TABLE IF NOT EXISTS budgets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            family_id INTEGER NOT NULL,
            category TEXT NOT NULL,
            limit_amount REAL NOT NULL,
            month TEXT NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(family_id, category, month),
            FOREIGN KEY (family_id) REFERENCES families(id)
        );
    ''')
    con.commit()
    con.close()


# ==================== HELPERS ====================
def gen_invite_code():
    alphabet = 'ABCDEFGHJKLMNPQRSTUVWXYZ23456789'
    return ''.join(secrets.choice(alphabet) for _ in range(8))


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get('user_id'):
            return jsonify({'error': 'unauthorized'}), 401
        return f(*args, **kwargs)
    return wrapper


def current_user():
    uid = session.get('user_id')
    if not uid:
        return None
    return get_db().execute(
        'SELECT u.*, f.name AS family_name, f.invite_code '
        'FROM users u LEFT JOIN families f ON u.family_id = f.id '
        'WHERE u.id = ?', (uid,)
    ).fetchone()


# ==================== PAGES ====================
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/healthz')
def healthz():
    return 'ok'


# ==================== AUTH ====================
@app.route('/api/me')
def api_me():
    user = current_user()
    if not user:
        return jsonify({'authenticated': False})
    return jsonify({
        'authenticated': True,
        'user': {
            'id': user['id'],
            'username': user['username'],
            'display_name': user['display_name'],
            'family_id': user['family_id'],
            'family_name': user['family_name'],
            'invite_code': user['invite_code'],
        }
    })


@app.route('/api/register', methods=['POST'])
def api_register():
    data = request.get_json() or {}
    username = (data.get('username') or '').strip().lower()
    password = data.get('password') or ''
    display_name = (data.get('display_name') or '').strip()
    family_action = data.get('family_action')  # 'create' | 'join'
    family_name = (data.get('family_name') or '').strip()
    invite_code = (data.get('invite_code') or '').strip().upper()

    if not username or len(username) < 3:
        return jsonify({'error': 'Логин минимум 3 символа'}), 400
    if len(password) < 6:
        return jsonify({'error': 'Пароль минимум 6 символов'}), 400
    if not display_name:
        return jsonify({'error': 'Укажите ваше имя'}), 400
    if family_action not in ('create', 'join'):
        return jsonify({'error': 'Выберите: создать или присоединиться'}), 400

    db = get_db()
    if db.execute('SELECT 1 FROM users WHERE username=?', (username,)).fetchone():
        return jsonify({'error': 'Такой логин уже занят'}), 400

    now = datetime.utcnow().isoformat()

    if family_action == 'create':
        if not family_name:
            return jsonify({'error': 'Введите название семьи'}), 400
        code = gen_invite_code()
        while db.execute('SELECT 1 FROM families WHERE invite_code=?', (code,)).fetchone():
            code = gen_invite_code()
        cur = db.execute(
            'INSERT INTO families (name, invite_code, created_at) VALUES (?, ?, ?)',
            (family_name, code, now)
        )
        family_id = cur.lastrowid
    else:
        if not invite_code:
            return jsonify({'error': 'Введите код семьи'}), 400
        fam = db.execute('SELECT id FROM families WHERE invite_code=?', (invite_code,)).fetchone()
        if not fam:
            return jsonify({'error': 'Семья с таким кодом не найдена'}), 400
        family_id = fam['id']

    db.execute(
        'INSERT INTO users (username, password_hash, display_name, family_id, created_at) '
        'VALUES (?, ?, ?, ?, ?)',
        (username, generate_password_hash(password), display_name, family_id, now)
    )
    db.commit()

    user = db.execute('SELECT id FROM users WHERE username=?', (username,)).fetchone()
    session.permanent = True
    session['user_id'] = user['id']
    return jsonify({'ok': True})


@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.get_json() or {}
    username = (data.get('username') or '').strip().lower()
    password = data.get('password') or ''
    db = get_db()
    user = db.execute('SELECT * FROM users WHERE username=?', (username,)).fetchone()
    if not user or not check_password_hash(user['password_hash'], password):
        return jsonify({'error': 'Неверный логин или пароль'}), 401
    session.permanent = True
    session['user_id'] = user['id']
    return jsonify({'ok': True})


@app.route('/api/logout', methods=['POST'])
def api_logout():
    session.clear()
    return jsonify({'ok': True})


# ==================== FAMILY ====================
@app.route('/api/family/members')
@login_required
def api_family_members():
    user = current_user()
    db = get_db()
    rows = db.execute(
        'SELECT id, display_name, username FROM users WHERE family_id=? ORDER BY display_name',
        (user['family_id'],)
    ).fetchall()
    return jsonify([dict(r) for r in rows])


# ==================== EXPENSES ====================
@app.route('/api/expenses', methods=['GET'])
@login_required
def api_list_expenses():
    user = current_user()
    month = request.args.get('month')  # YYYY-MM
    t = request.args.get('type')  # expense | income | all
    db = get_db()

    q = ('SELECT e.*, u.display_name AS who '
         'FROM expenses e JOIN users u ON e.user_id = u.id '
         'WHERE e.family_id = ?')
    params = [user['family_id']]
    if month:
        q += ' AND e.date LIKE ?'
        params.append(month + '%')
    if t in ('expense', 'income'):
        q += ' AND e.type = ?'
        params.append(t)
    q += ' ORDER BY e.date DESC, e.id DESC'
    rows = db.execute(q, params).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route('/api/expenses', methods=['POST'])
@login_required
def api_add_expense():
    user = current_user()
    data = request.get_json() or {}
    t = data.get('type', 'expense')
    if t not in ('expense', 'income'):
        return jsonify({'error': 'Неверный тип'}), 400
    name = (data.get('name') or '').strip()
    category = (data.get('category') or '').strip()
    date = (data.get('date') or '').strip()
    try:
        amount = float(data.get('amount'))
        if amount <= 0:
            raise ValueError
    except (TypeError, ValueError):
        return jsonify({'error': 'Неверная сумма'}), 400
    if not name or not category or not date:
        return jsonify({'error': 'Заполните все поля'}), 400

    db = get_db()
    cur = db.execute(
        'INSERT INTO expenses (family_id, user_id, type, name, amount, category, date, created_at) '
        'VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
        (user['family_id'], user['id'], t, name[:200], amount,
         category[:50], date[:10], datetime.utcnow().isoformat())
    )
    db.commit()
    row = db.execute(
        'SELECT e.*, u.display_name AS who FROM expenses e '
        'JOIN users u ON e.user_id = u.id WHERE e.id=?', (cur.lastrowid,)
    ).fetchone()
    result = dict(row)
    socketio.emit('refresh', {'reason': 'expense_added'},
                  room=f'family_{user["family_id"]}')
    return jsonify(result), 201


@app.route('/api/expenses/<int:eid>', methods=['DELETE'])
@login_required
def api_del_expense(eid):
    user = current_user()
    db = get_db()
    db.execute('DELETE FROM expenses WHERE id=? AND family_id=?',
               (eid, user['family_id']))
    db.commit()
    socketio.emit('refresh', {'reason': 'expense_deleted'},
                  room=f'family_{user["family_id"]}')
    return jsonify({'ok': True})


# ==================== MONTHS ====================
@app.route('/api/months')
@login_required
def api_months():
    user = current_user()
    db = get_db()
    rows = db.execute(
        "SELECT DISTINCT substr(date,1,7) AS m FROM expenses "
        "WHERE family_id=? ORDER BY m DESC", (user['family_id'],)
    ).fetchall()
    return jsonify([r['m'] for r in rows])


# ==================== BUDGETS ====================
@app.route('/api/budgets', methods=['GET'])
@login_required
def api_list_budgets():
    user = current_user()
    month = request.args.get('month', datetime.utcnow().strftime('%Y-%m'))
    db = get_db()
    rows = db.execute(
        'SELECT * FROM budgets WHERE family_id=? AND month=? ORDER BY category',
        (user['family_id'], month)
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route('/api/budgets', methods=['POST'])
@login_required
def api_set_budget():
    user = current_user()
    data = request.get_json() or {}
    category = (data.get('category') or '').strip()
    month = (data.get('month') or '').strip()
    try:
        limit = float(data.get('limit_amount'))
        if limit <= 0:
            raise ValueError
    except (TypeError, ValueError):
        return jsonify({'error': 'Неверный лимит'}), 400
    if not category or not month:
        return jsonify({'error': 'Заполните поля'}), 400

    db = get_db()
    db.execute(
        'INSERT INTO budgets (family_id, category, limit_amount, month, created_at) '
        'VALUES (?, ?, ?, ?, ?) '
        'ON CONFLICT(family_id, category, month) DO UPDATE SET limit_amount=excluded.limit_amount',
        (user['family_id'], category[:50], limit, month[:7],
         datetime.utcnow().isoformat())
    )
    db.commit()
    socketio.emit('refresh', {'reason': 'budget_changed'},
                  room=f'family_{user["family_id"]}')
    return jsonify({'ok': True})


@app.route('/api/budgets/<int:bid>', methods=['DELETE'])
@login_required
def api_del_budget(bid):
    user = current_user()
    db = get_db()
    db.execute('DELETE FROM budgets WHERE id=? AND family_id=?',
               (bid, user['family_id']))
    db.commit()
    socketio.emit('refresh', {'reason': 'budget_deleted'},
                  room=f'family_{user["family_id"]}')
    return jsonify({'ok': True})


# ==================== STATS ====================
@app.route('/api/stats')
@login_required
def api_stats():
    user = current_user()
    month = request.args.get('month', datetime.utcnow().strftime('%Y-%m'))
    db = get_db()
    fam = user['family_id']

    total_exp = db.execute(
        "SELECT COALESCE(SUM(amount),0) AS s FROM expenses "
        "WHERE family_id=? AND type='expense' AND date LIKE ?",
        (fam, month + '%')
    ).fetchone()['s']
    total_inc = db.execute(
        "SELECT COALESCE(SUM(amount),0) AS s FROM expenses "
        "WHERE family_id=? AND type='income' AND date LIKE ?",
        (fam, month + '%')
    ).fetchone()['s']

    by_cat = [dict(r) for r in db.execute(
        "SELECT category, SUM(amount) AS total FROM expenses "
        "WHERE family_id=? AND type='expense' AND date LIKE ? "
        "GROUP BY category ORDER BY total DESC",
        (fam, month + '%')
    ).fetchall()]

    by_day = [dict(r) for r in db.execute(
        "SELECT date, SUM(amount) AS total FROM expenses "
        "WHERE family_id=? AND type='expense' AND date LIKE ? "
        "GROUP BY date ORDER BY date",
        (fam, month + '%')
    ).fetchall()]

    by_user = [dict(r) for r in db.execute(
        "SELECT u.display_name AS who, e.type, SUM(e.amount) AS total "
        "FROM expenses e JOIN users u ON e.user_id = u.id "
        "WHERE e.family_id=? AND e.date LIKE ? "
        "GROUP BY u.display_name, e.type ORDER BY total DESC",
        (fam, month + '%')
    ).fetchall()]

    budgets = [dict(r) for r in db.execute(
        'SELECT * FROM budgets WHERE family_id=? AND month=?',
        (fam, month)
    ).fetchall()]

    # прогресс по бюджетам
    cat_totals = {r['category']: r['total'] for r in by_cat}
    budget_status = []
    for b in budgets:
        spent = cat_totals.get(b['category'], 0)
        pct = (spent / b['limit_amount'] * 100) if b['limit_amount'] else 0
        budget_status.append({
            'id': b['id'],
            'category': b['category'],
            'limit': b['limit_amount'],
            'spent': spent,
            'percent': round(pct, 1),
            'over': spent > b['limit_amount'],
        })

    # сравнение с прошлым месяцем
    y, m = map(int, month.split('-'))
    pm = f"{y-1:04d}-12" if m == 1 else f"{y:04d}-{m-1:02d}"
    prev_exp = db.execute(
        "SELECT COALESCE(SUM(amount),0) AS s FROM expenses "
        "WHERE family_id=? AND type='expense' AND date LIKE ?",
        (fam, pm + '%')
    ).fetchone()['s']

    return jsonify({
        'month': month,
        'total_expense': total_exp,
        'total_income': total_inc,
        'balance': total_inc - total_exp,
        'by_category': by_cat,
        'by_day': by_day,
        'by_user': by_user,
        'budgets': budget_status,
        'prev_month_expense': prev_exp,
    })


# ==================== EXPORT EXCEL ====================
@app.route('/api/export/excel')
@login_required
def api_export_excel():
    user = current_user()
    month = request.args.get('month', datetime.utcnow().strftime('%Y-%m'))
    db = get_db()

    rows = db.execute(
        'SELECT e.date, e.type, e.category, e.name, e.amount, u.display_name AS who '
        'FROM expenses e JOIN users u ON e.user_id = u.id '
        'WHERE e.family_id=? AND e.date LIKE ? '
        'ORDER BY e.date, e.id',
        (user['family_id'], month + '%')
    ).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = f'Расходы {month}'

    headers = ['Дата', 'Тип', 'Категория', 'Наименование', 'Сумма', 'Кто внёс']
    ws.append(headers)
    hdr_font = Font(bold=True, color='FFFFFF')
    hdr_fill = PatternFill('solid', fgColor='667EEA')
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center')

    total_exp = total_inc = 0
    for r in rows:
        ws.append([
            r['date'],
            'Расход' if r['type'] == 'expense' else 'Доход',
            r['category'],
            r['name'],
            r['amount'],
            r['who'],
        ])
        if r['type'] == 'expense':
            total_exp += r['amount']
        else:
            total_inc += r['amount']

    ws.append([])
    ws.append(['', '', '', 'Итого расходов', total_exp, ''])
    ws.append(['', '', '', 'Итого доходов', total_inc, ''])
    ws.append(['', '', '', 'Баланс', total_inc - total_exp, ''])

    for col, width in zip('ABCDEF', [12, 10, 18, 30, 12, 15]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f'budget_{month}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


# ==================== SOCKETIO ====================
@socketio.on('connect')
def sio_connect():
    uid = session.get('user_id')
    if not uid:
        return False
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    u = db.execute('SELECT family_id FROM users WHERE id=?', (uid,)).fetchone()
    db.close()
    if not u or not u['family_id']:
        return False
    join_room(f'family_{u["family_id"]}')
    emit('connected', {'family_id': u['family_id']})


# ==================== BOOT ====================
init_db()

if __name__ == '__main__':
    socketio.run(app, host='0.0.0.0',
                 port=int(os.environ.get('PORT', 5000)),
                 debug=True, allow_unsafe_werkzeug=True)